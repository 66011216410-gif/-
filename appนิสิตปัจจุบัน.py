import io
import re
from copy import copy
from datetime import datetime

import pandas as pd
import streamlit as st

st.set_page_config(page_title="ระบบสถิตินิสิตบัณฑิตศึกษา", page_icon="📊", layout="wide")



REQUIRED_COLUMNS = [
    "ปีที่เข้า", "ภาคการศึกษาที่เข้า", "รหัสนิสิต", "คณะ", "วิทยาเขต", "สาขา",
    "ระดับ", "รหัสสถานะนิสิต", "สถานะนิสิต", "ปีที่จบ", "เทอมที่จบ", "วันที่จบ"
]

# สถานะที่ต้องไม่นำมาคำนวณสถิติ 4 ช่องใหม่
STATUS_EXCLUDE = ["พ้นสภาพ (เสียชีวิต)", "พ้นสภาพ", "ลาออก"]

STATUS_DISPLAY = [
    "นิสิตปัจจุบัน", "พ้นสภาพ (คณบดีอนุมัติ)", "พ้นสภาพ (เสียชีวิต)",
    "พ้นสภาพ", "รักษาสภาพนิสิต", "ลาพักการเรียน", "ลาออก"
]


def clean_text(x):
    if pd.isna(x):
        return ""
    return str(x).strip()


def num(x):
    try:
        return float(x)
    except Exception:
        return None


def calc_duration(row):
    """1 เทอม = 0.5 ปี และภาคเดียวกันของปีเดียวกัน = 0.5 ปี"""
    y0, s0 = num(row.get("ปีที่เข้า")), num(row.get("ภาคการศึกษาที่เข้า"))
    y1, s1 = num(row.get("ปีที่จบ")), num(row.get("เทอมที่จบ"))
    if None in (y0, s0, y1, s1):
        return None
    semesters = (y1 - y0) * 2 + (s1 - s0) + 1
    d = semesters * 0.5
    return round(d, 1) if d > 0 else None


def normalize_level(x):
    x = clean_text(x)
    if "เอก" in x:
        return "ป.เอก"
    if "โท" in x:
        return "ป.โท"
    return x or "ไม่ระบุ"


def status_group(x):
    s = clean_text(x)

    # กลุ่ม "พ้นสภาพ" ต้องตรวจสอบก่อนสถานะอื่น
    # เพราะบางข้อความมีคำว่า "สำเร็จการศึกษา", "รักษาสภาพ" หรือ "ลาพัก" อยู่ภายใน
    if (
        s == "นิสิตที่ไม่มารายงานตัว"
        or any(
            code in s
            for code in [
                "38.4 ",
                "38.4.1 ",
                "38.4.12 ",
                "38.4.14 ",
                "38.4.2 ",
                "38.4.3 ",
                "38.4.4 ",
                "38.4.5 ",
                "38.4.6 ",
                "38.4.8 ",
                "38.6 ",
                "38.7 ",
            ]
        )
    ):
        return "พ้นสภาพ"

    # รวม 37.3 และ 38.3 เป็น "พ้นสภาพ (คณบดีอนุมัติ)"
    if (
        "37.3 ลาออกโดยได้รับอนุมัติจากคณบดีบัณฑิตวิทยาลัย" in s
        or "38.3 ลาออกโดยได้รับอนุมัติจากคณบดีบัณฑิตวิทยาลัย" in s
    ):
        return "พ้นสภาพ (คณบดีอนุมัติ)"

    # รวม 19.1 และ 37.2 เป็น "พ้นสภาพ (เสียชีวิต)"
    if "19.1 ตาย" in s or "37.2 ตาย" in s:
        return "พ้นสภาพ (เสียชีวิต)"
    if "เสียชีวิต" in s:
        return "พ้นสภาพ (เสียชีวิต)"

    if "พ้นสภาพคืนไม่ได้" in s:
        return "พ้นสภาพคืนไม่ได้"
    if "สำเร็จการศึกษา" in s:
        return "สำเร็จการศึกษา"
    if "สภาพสมบูรณ์" in s:
        return "สภาพสมบูรณ์"
    if "รักษาสภาพ" in s:
        return "รักษาสภาพนิสิต"
    if "ลาพัก" in s:
        return "ลาพักการเรียน"
    if "ลาออก" in s:
        return "ลาออก"
    if "คณบดี" in s:
        return "พ้นสภาพ (คณบดีอนุมัติ)"
    if "พ้นสภาพ" in s:
        return "พ้นสภาพ"
    if "ปัจจุบัน" in s or "กำลังศึกษา" in s:
        return "นิสิตปัจจุบัน"
    return s or "ไม่ระบุ"


def cut_plan_type(x):
    s = clean_text(x)
    if not s:
        return ""
    return re.split(
        r"\s*(?:แผน|แบบ)(?:\s|[:：\-/]|$).*$",
        s,
        maxsplit=1,
    )[0].strip(" -:：/|")


def read_excel(uploaded):
    # ไม่บังคับชื่อ Sheet: ค้นหา Sheet ใดก็ได้ที่มีคอลัมน์ข้อมูลนิสิตครบ
    sheets = pd.read_excel(uploaded, sheet_name=None)
    for sheet_name, sheet_df in sheets.items():
        sheet_df.columns = [clean_text(c) for c in sheet_df.columns]
        if all(c in sheet_df.columns for c in REQUIRED_COLUMNS):
            return sheet_df

    # หากไม่พบ Sheet ที่ตรง ให้แจ้งคอลัมน์ที่ต้องมี
    available = []
    for sheet_name, sheet_df in sheets.items():
        available.append(f"{sheet_name}: {', '.join(map(str, sheet_df.columns))}")
    raise ValueError(
        "ไม่พบ Sheet ที่มีคอลัมน์ข้อมูลนิสิตครบ ระบบไม่บังคับชื่อ Sheet "
        "แต่ต้องมีคอลัมน์: " + ", ".join(REQUIRED_COLUMNS)
    )


def metric_row(label, g, limit, durations):
    r = {"ปีที่เข้า": label, "จำนวนนิสิตรับเข้า(คน)": len(g)}

    for d in durations:
        r[d] = int((g["ระยะเวลา(ปี)"] == d).sum())

    # จำนวนนิสิตจบทั้งหมด = นับจากคอลัมน์ "สถานะนิสิต" โดยตรง
    # ต้องเป็นค่า "สำเร็จการศึกษา" เท่านั้น ไม่รวมสถานะอื่น
    grads = g[g["สถานะนิสิต"].map(clean_text) == "สำเร็จการศึกษา"].copy()
    valid = g[~g["สถานะกลุ่ม"].isin(STATUS_EXCLUDE)]
    valid_duration = valid["ระยะเวลา(ปี)"].dropna()

    # สถิติเดิม
    # นับเป็น "จำนวนนิสิต" โดยนับรหัสนิสิตไม่ซ้ำ
    # เพื่อไม่ให้นิสิตคนเดียวที่มีหลายรายการถูกนับซ้ำ
    if "รหัสนิสิต" in grads.columns:
        r["จำนวนนิสิตจบ_ทั้งหมด"] = grads["รหัสนิสิต"].map(clean_text).replace("", pd.NA).dropna().nunique()
        on_time = grads[grads["ระยะเวลา(ปี)"] <= limit]
        r["จำนวนนิสิตจบ_ตามหลักสูตร"] = on_time["รหัสนิสิต"].map(clean_text).replace("", pd.NA).dropna().nunique()
    else:
        r["จำนวนนิสิตจบ_ทั้งหมด"] = len(grads)
        r["จำนวนนิสิตจบ_ตามหลักสูตร"] = int(
            (grads["ระยะเวลา(ปี)"] <= limit).sum()
        )

    # สถิติ 4 ช่องใหม่
    # สูตรที่กำหนด:
    # จำนวนนิสิตที่ไม่นับสถานะ
    # = จำนวนนิสิตรับเข้า - พ้นสภาพ(เสียชีวิต) - พ้นสภาพ - ลาออก - พ้นสภาพ (คณบดีอนุมัติ)
    death_mask = g["สถานะกลุ่ม"] == "พ้นสภาพ (เสียชีวิต)"
    dropout_mask = g["สถานะกลุ่ม"] == "พ้นสภาพ"
    resign_mask = g["สถานะกลุ่ม"] == "ลาออก"
    dean_mask = g["สถานะกลุ่ม"] == "พ้นสภาพ (คณบดีอนุมัติ)"
    excluded_mask = death_mask | dropout_mask | resign_mask | dean_mask
    excluded_count = int(excluded_mask.sum())
    r["จำนวนนิสิตที่ไม่นับสถานะ"] = len(g) - excluded_count
    valid = g[~excluded_mask]
    valid_duration = valid["ระยะเวลา(ปี)"].dropna()
    r["เฉลี่ยระยะเวลาที่ใช้(ปี) ไม่นับรวมนิสิตสถานะ เสียชีวิต พ้นสภาพ ลาออก"] = (
        round(valid_duration.mean(), 2) if len(valid_duration) else 0
    )
    r["เฉลี่ยระยะเวลาที่ใช้(ปี)"] = (
        round(grads["ระยะเวลา(ปี)"].dropna().mean(), 2)
        if grads["ระยะเวลา(ปี)"].notna().any()
        else 0
    )
    r["%จบตามเวลา"] = (
        r["จำนวนนิสิตจบ_ตามหลักสูตร"] * 100 / r["จำนวนนิสิตที่ไม่นับสถานะ"]
        if r["จำนวนนิสิตที่ไม่นับสถานะ"]
        else 0
    )

    r["ยังไม่จบ_ทั้งหมด"] = len(g) - len(grads)

    # คอลัมน์ "นิสิตปัจจุบัน" ให้นับเฉพาะสถานะ
    # "นิสิตปัจจุบัน สภาพสมบูรณ์" เท่านั้น
    r["นิสิตปัจจุบัน"] = int(
        (g["สถานะนิสิต"].map(clean_text) == "นิสิตปัจจุบัน สภาพสมบูรณ์").sum()
    )

    for stt in STATUS_DISPLAY:
        if stt == "นิสิตปัจจุบัน":
            continue
        r[stt] = int((g["สถานะกลุ่ม"] == stt).sum())

    r["พ้นสภาพคืนไม่ได้"] = int(
        (g["สถานะกลุ่ม"] == "พ้นสภาพคืนไม่ได้").sum()
    )
    return r


def build_stats(df):
    work = df.copy()
    work["ระดับ"] = work["ระดับ"].map(normalize_level)
    work["สถานะกลุ่ม"] = work["สถานะนิสิต"].map(status_group)

    # ผู้ที่ไม่ใช่ผู้สำเร็จการศึกษา จะไม่มีข้อมูลปี/เทอม/วันที่จบในการคำนวณระยะเวลา
    non_graduated = work["สถานะนิสิต"].map(clean_text) != "สำเร็จการศึกษา"
    work.loc[
        non_graduated, ["ปีที่จบ", "เทอมที่จบ", "วันที่จบ"]
    ] = pd.NA

    work["ระยะเวลา(ปี)"] = work.apply(calc_duration, axis=1)
    work["สาขาสถิติ"] = work["สาขา"].map(cut_plan_type)

    durations = [x / 2 for x in range(1, 23)]
    rows = []

    levels = [x for x in ["ป.โท", "ป.เอก"] if x in work["ระดับ"].unique()]
    levels += [x for x in work["ระดับ"].unique() if x not in levels]

    for level in levels:
        g = work[work["ระดับ"] == level]
        if g.empty:
            continue

        # ป.โท 2 ปี / ป.เอก 4 ปี
        limit = 2 if level == "ป.โท" else 4
        rows.append(metric_row(level, g, limit, durations))

        years = sorted(
            pd.to_numeric(g["ปีที่เข้า"], errors="coerce").dropna().unique()
        )
        for year in years:
            gy = g[
                pd.to_numeric(g["ปีที่เข้า"], errors="coerce") == year
            ]
            rows.append(metric_row(int(year), gy, limit, durations))

            for faculty, gf in gy.groupby("คณะ", dropna=False, sort=True):
                faculty = clean_text(faculty)
                if not faculty:
                    continue

                rows.append(metric_row("คณะ" + faculty, gf, limit, durations))

                for program, gp in gf.groupby(
                    "สาขาสถิติ", dropna=False, sort=True
                ):
                    program = clean_text(program)
                    if not program:
                        continue
                    rows.append(metric_row(program, gp, limit, durations))

    # แถวล่างสุด "รวมทั้งหมด" รวมข้อมูลทุกระดับ/ทุกปี
    if rows:
        all_g = work.copy()
        total = metric_row("รวมทั้งหมด", all_g, 2, durations)

        # ผู้สำเร็จการศึกษาตามหลักสูตร: ป.โท <= 2 ปี และ ป.เอก <= 4 ปี
        grads_all = all_g[all_g["สถานะนิสิต"].map(clean_text) == "สำเร็จการศึกษา"].copy()
        if "รหัสนิสิต" in grads_all.columns:
            total["จำนวนนิสิตจบ_ทั้งหมด"] = grads_all["รหัสนิสิต"].map(clean_text).replace("", pd.NA).dropna().nunique()
            on_time = (
                ((grads_all["ระดับ"] == "ป.โท") & (grads_all["ระยะเวลา(ปี)"] <= 2))
                | ((grads_all["ระดับ"] == "ป.เอก") & (grads_all["ระยะเวลา(ปี)"] <= 4))
            )
            total["จำนวนนิสิตจบ_ตามหลักสูตร"] = grads_all.loc[on_time, "รหัสนิสิต"].map(clean_text).replace("", pd.NA).dropna().nunique()
        total["%จบตามเวลา"] = (
            total["จำนวนนิสิตจบ_ตามหลักสูตร"] * 100 / total["จำนวนนิสิตที่ไม่นับสถานะ"]
            if total["จำนวนนิสิตที่ไม่นับสถานะ"] else 0
        )
        rows.append(total)

    # 40 คอลัมน์ตรงกับ Sheet "สถิติ" เดิม
    # ใช้ชื่อภายในที่ไม่ซ้ำกันสำหรับ % เดิม เพื่อป้องกัน pyarrow/Streamlit
    columns = ["ปีที่เข้า", "จำนวนนิสิตรับเข้า(คน)"] + durations + [
        "จำนวนนิสิตจบ_ทั้งหมด",
        "จำนวนนิสิตจบ_ตามหลักสูตร",
        "%จบตามเวลา_เดิม",
        "ยังไม่จบ_ทั้งหมด",
        "นิสิตปัจจุบัน",
        "พ้นสภาพ (คณบดีอนุมัติ)",
        "พ้นสภาพ (เสียชีวิต)",
        "พ้นสภาพ",
        "รักษาสภาพนิสิต",
        "ลาพักการเรียน",
        "ลาออก",
        "เฉลี่ยระยะเวลาที่ใช้(ปี)",
        "เฉลี่ยระยะเวลาที่ใช้(ปี) ไม่นับรวมนิสิตสถานะ เสียชีวิต พ้นสภาพ ลาออก",
        "จำนวนนิสิตที่ไม่นับสถานะ เสียชีวิต พ้นสภาพ ลาออก พ้นสภาพ (คณบดีอนุมัติ)",
        "ตามระยะเวลาของหลักสูตร 2 ปี/ 4 ปี (คน)",
        "%จบตามเวลา",
    ]

    out = []
    for r in rows:
        vals = [r["ปีที่เข้า"], r["จำนวนนิสิตรับเข้า(คน)"]]
        vals += [r[d] for d in durations]
        # คอลัมน์ Y:AN
        # AL = รับเข้า - คณบดีอนุมัติ - เสียชีวิต - พ้นสภาพ - ลาออก
        # AM = จำนวนจบตามระยะเวลาหลักสูตร
        # AN = % จบตามเวลา
        excluded_count = (
            r["พ้นสภาพ (คณบดีอนุมัติ)"]
            + r["พ้นสภาพ (เสียชีวิต)"]
            + r["พ้นสภาพ"]
            + r["ลาออก"]
        )
        not_counted = r["จำนวนนิสิตรับเข้า(คน)"] - excluded_count
        on_time_count = r["จำนวนนิสิตจบ_ตามหลักสูตร"]
        on_time_pct = (
            on_time_count * 100 / not_counted
            if not_counted else 0
        )

        # บังคับค่าของ 3 คอลัมน์ท้ายให้ถูกต้องก่อนสร้าง DataFrame
        r["จำนวนนิสิตที่ไม่นับสถานะ"] = not_counted
        r["จำนวนนิสิตจบ_ตามหลักสูตร"] = on_time_count
        r["%จบตามเวลา"] = on_time_pct

        vals += [
            r["จำนวนนิสิตจบ_ทั้งหมด"],
            r["จำนวนนิสิตจบ_ตามหลักสูตร"],
            r["%จบตามเวลา"],
            r["ยังไม่จบ_ทั้งหมด"],
            r["นิสิตปัจจุบัน"],
            r["พ้นสภาพ (คณบดีอนุมัติ)"],
            r["พ้นสภาพ (เสียชีวิต)"],
            r["พ้นสภาพ"],
            r["รักษาสภาพนิสิต"],
            r["ลาพักการเรียน"],
            r["ลาออก"],
            r["เฉลี่ยระยะเวลาที่ใช้(ปี)"],
            r["เฉลี่ยระยะเวลาที่ใช้(ปี) ไม่นับรวมนิสิตสถานะ เสียชีวิต พ้นสภาพ ลาออก"],
            not_counted,
            on_time_count,
            on_time_pct,
        ]
        out.append(vals)

    return work, pd.DataFrame(out, columns=columns)


def copy_row_style(ws, source_row, target_row, max_col=40):
    if source_row == target_row:
        return

    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

    for c in range(1, max_col + 1):
        src = ws.cell(source_row, c)
        dst = ws.cell(target_row, c)

        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.protection:
            dst.protection = copy(src.protection)


def make_excel(original_uploaded, original, processed, stats):
    """สร้าง Sheet สถิติให้เป็น Template ตามแบบที่กำหนด พร้อมคงสี/เส้น/รูปแบบ"""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    original_uploaded.seek(0)
    wb = load_workbook(original_uploaded)

    if "ข้อมูลประมวลผล" in wb.sheetnames:
        del wb["ข้อมูลประมวลผล"]

    ws_processed = wb.create_sheet("ข้อมูลประมวลผล")
    ws_processed.append(processed.columns.tolist())
    for row in processed.where(pd.notna(processed), None).values.tolist():
        ws_processed.append(row)

    ws_processed.freeze_panes = "A2"
    ws_processed.auto_filter.ref = ws_processed.dimensions

    if "สถิติ" in wb.sheetnames:
        ws = wb["สถิติ"]
    else:
        ws = wb.create_sheet("สถิติ")
    # ============================================================    # TEMPLATE SHEET "สถิติ"
    # ============================================================
    gray = PatternFill(fill_type="solid", fgColor="808080")
    red_fill = PatternFill(fill_type="solid", fgColor="FF0000")
    pink_fill = PatternFill(fill_type="solid", fgColor="FF99CC")
    green_fill = PatternFill(fill_type="solid", fgColor="A9D18E")
    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    blue_fill = PatternFill(fill_type="solid", fgColor="2F5597")

    white_font = Font(name="Tahoma", size=10, color="FFFFFF", bold=True)
    black_font = Font(name="Tahoma", size=10, color="000000")
    red_font = Font(name="Tahoma", size=10, color="FF0000", bold=True)
    blue_font = Font(name="Tahoma", size=10, color="FFFFFF", bold=True)

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    center = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    # ยกเลิก merge เดิมเฉพาะบริเวณหัวตาราง
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row <= 3 and merged.min_col <= 40:
            ws.unmerge_cells(str(merged))

    # ล้างหัวตาราง A1:AN3 แล้วสร้างโครงใหม่
    for r in range(1, 4):
        for c in range(1, 41):
            ws.cell(r, c).value = None
            ws.cell(r, c).fill = gray
            ws.cell(r, c).font = white_font
            ws.cell(r, c).alignment = center
            ws.cell(r, c).border = border

    # กลุ่มหัวตาราง
    ws.merge_cells("A1:A3")
    ws.merge_cells("B1:B3")
    ws.merge_cells("C1:X1")
    ws.merge_cells("Y1:AA1")
    ws.merge_cells("AB1:AI1")

    for col in range(3, 25):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

    for col in range(25, 28):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

    for col in range(28, 36):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

    for col in range(36, 41):
        ws.merge_cells(start_row=1, start_column=col, end_row=3, end_column=col)

    ws["A1"] = "ปีที่เข้า"
    ws["B1"] = "จำนวน\nนิสิต\nรับเข้า(คน)"
    ws["C1"] = "ระยะเวลา(ปี)"
    ws["Y1"] = "จำนวนนิสิตจบ"
    ws["AB1"] = "จำนวนนิสิตที่ยังไม่จบ"

    durations = [x / 2 for x in range(1, 23)]
    for i, value in enumerate(durations, start=3):
        ws.cell(2, i).value = value

    ws["Y2"] = "ทั้งหมด(คน)"
    ws["Z2"] = "ตาม\nระยะเวลา\nของหลักสูตร\n2 ปี/ 4 ปี\n(คน)"
    ws["AA2"] = "% จบ\nตามเวลา"

    ws["AB2"] = "ทั้งหมด"
    ws["AC2"] = "นิสิตปัจจุบัน"
    ws["AD2"] = "พ้นสภาพ (คณบดีอนุมัติ)"
    ws["AE2"] = "พ้นสภาพ (เสียชีวิต)"
    ws["AF2"] = "พ้นสภาพ"
    ws["AG2"] = "รักษาสภาพนิสิต"
    ws["AH2"] = "ลาพักการเรียน"
    ws["AI2"] = "ลาออก"

    ws["AJ1"] = "เฉลี่ย\nระยะเวลาที่ใช้(ปี)"
    ws["AK1"] = (
        "เฉลี่ยระยะเวลาที่ใช้(ปี)\n"
        "ไม่นับรวมนิสิตสถานะ\n"
        "เสียชีวิต พ้นสภาพ ลาออก"
    )
    ws["AL1"] = (
        "จำนวนนิสิตที่ไม่นับสถานะ\n"
        "เสียชีวิต พ้นสภาพ ลาออก"
    )
    ws["AM1"] = "ตาม\nระยะเวลาของหลักสูตร\n2 ปี/ 4 ปี\n(คน)"
    ws["AN1"] = "% จบ\nตามเวลา"

    # รูปแบบหัวตาราง
    for r in range(1, 4):
        for c in range(1, 41):
            cell = ws.cell(r, c)
            cell.fill = gray
            cell.font = white_font
            cell.alignment = center
            cell.border = border

    # คอลัมน์ AL ตาม Template ใช้สีน้ำเงิน
    ws["AL1"].fill = blue_fill
    ws["AL1"].font = blue_font

    # หัวข้อสถานะที่เป็นสีแดงตาม Template
    for cell_ref in ["AD2", "AE2", "AF2", "AI2"]:
        ws[cell_ref].font = red_font

    # ความกว้างคอลัมน์
    widths = {
        "A": 11, "B": 12,
        "Y": 12, "Z": 13, "AA": 10,
        "AB": 11, "AC": 12, "AD": 15, "AE": 15,
        "AF": 11, "AG": 13, "AH": 13, "AI": 10,
        "AJ": 12, "AK": 16, "AL": 15, "AM": 13, "AN": 10,
    }
    for col in range(3, 25):
        widths[chr(64 + col)] = 7
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 27
    ws.row_dimensions[3].height = 27

    # ============================================================
    # DATA ROWS — เริ่มที่แถว 4 ตาม Template
    # ============================================================
    first_data_row = 4
    old_last = ws.max_row

    # ล้างค่าเดิม แต่คงโครงสร้างหัวตาราง
    if old_last >= first_data_row:
        for r in range(first_data_row, old_last + 1):
            for c in range(1, 41):
                ws.cell(r, c).value = None

    # ถ้าจำนวนแถวใหม่มากกว่าเดิม ให้เพิ่มแถว
    needed_last = first_data_row + len(stats) - 1
    if needed_last > old_last:
        ws.insert_rows(old_last + 1, needed_last - old_last)

    # เขียนข้อมูล 40 คอลัมน์
    for i, (_, row) in enumerate(stats.iterrows(), start=first_data_row):
        # อ่าน 37 คอลัมน์แรกก่อน แล้วคำนวณ 3 คอลัมน์ท้ายใหม่
        # เพื่อป้องกันค่า AL:AN หายจาก DataFrame/รูปแบบคอลัมน์
        values = [
            row.iloc[j] if pd.notna(row.iloc[j]) else None
            for j in range(37)
        ]

        # AL (38) = รับเข้า - คณบดีอนุมัติ - เสียชีวิต - พ้นสภาพ - ลาออก
        not_counted = (
            (values[1] or 0)
            - (values[29] or 0)
            - (values[30] or 0)
            - (values[31] or 0)
            - (values[34] or 0)
        )
        # AM (39) = จำนวนนิสิตจบตามระยะเวลาของหลักสูตร
        on_time = values[25] or 0
        # AN (40) = % จบตามเวลา
        on_time_pct = (on_time * 100 / not_counted) if not_counted else 0
        values += [not_counted, on_time, on_time_pct]

        label = clean_text(values[0])

        if label in ("ป.โท", "ป.เอก"):
            row_fill = red_fill
        elif re.fullmatch(r"\d{4}", label):
            row_fill = white_fill
        elif label.startswith("คณะ"):
            row_fill = pink_fill
        else:
            row_fill = green_fill

        for c, value in enumerate(values, start=1):
            cell = ws.cell(i, c)
            cell.value = value
            cell.fill = row_fill
            cell.font = black_font
            cell.alignment = center
            cell.border = border

        # แถวระดับใช้ตัวหนา
        if label in ("ป.โท", "ป.เอก"):
            for c in range(1, 41):
                ws.cell(i, c).font = Font(
                    name="Tahoma",
                    size=10,
                    color="000000",
                    bold=True,
                )

        # ยืนยันการเขียน 3 คอลัมน์ท้าย AL:AN โดยอ้างอิงชื่อคอลัมน์โดยตรง
        # ไม่พึ่งตำแหน่งของ DataFrame
        ws.cell(i, 38).value = row["จำนวนนิสิตที่ไม่นับสถานะ เสียชีวิต พ้นสภาพ ลาออก พ้นสภาพ (คณบดีอนุมัติ)"]
        ws.cell(i, 39).value = row["ตามระยะเวลาของหลักสูตร 2 ปี/ 4 ปี (คน)"]
        ws.cell(i, 40).value = row["%จบตามเวลา"]
        # ตัวเลขระยะเวลาเฉลี่ยและ % แสดง 2 ตำแหน่ง
        ws.cell(i, 36).number_format = "0.00"
        ws.cell(i, 37).number_format = "0.00"
        ws.cell(i, 40).number_format = "0.00"

    # ล้างแถวที่เกินจากข้อมูลใหม่
    for r in range(needed_last + 1, old_last + 1):
        for c in range(1, 41):
            ws.cell(r, c).value = None

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    if len(stats):
        ws.auto_filter.ref = f"A3:AN{needed_last}"

    # ============================================================
    # สร้าง Sheet "คณะ" และ "ปี" จาก Sheet "สถิติ" โดยตรง
    # ไม่คำนวณสถิติใหม่: ใช้ค่าจาก stats ที่คำนวณเสร็จแล้วทั้งหมด
    #
    # คณะ = คัดลอกสถิติ แล้วลบแถวระดับสาขาออก
    # ปี   = คัดลอกสถิติ แล้วลบแถวคณะและสาขาออก
    # ============================================================
    def write_copied_stats_sheet(sheet_name, source_stats, mode):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        sws = wb.copy_worksheet(ws)
        sws.title = sheet_name

        labels = source_stats["ปีที่เข้า"].map(clean_text)

        # แถวที่ต้องเก็บจาก Sheet "สถิติ"
        is_level = labels.isin(["ป.โท", "ป.เอก"])
        is_year = labels.str.match(r"^\d{4}(?:\.0)?$", na=False)
        is_faculty = labels.str.startswith("คณะ", na=False)
        is_total = labels == "รวมทั้งหมด"

        if mode == "faculty":
            # ป.โท -> ปี -> คณะ และ ป.เอก -> ปี -> คณะ
            keep = is_level | is_year | is_faculty | is_total
        else:
            # ป.โท -> ปี และ ป.เอก -> ปี
            keep = is_level | is_year | is_total

        copied = source_stats.loc[keep].copy().reset_index(drop=True)

        old_rows = sws.max_row
        # ข้อมูลของชีทสถิติเริ่มที่แถว 4 จึงต้องเขียนทับตั้งแต่แถว 4
        # เพื่อไม่ให้แถว ป.โท / ป.เอก จากชีทต้นแบบซ้ำกับข้อมูลที่คัดลอก
        for rr in range(4, old_rows + 1):
            for cc in range(1, 41):
                sws.cell(rr, cc).value = None
                # ล้างสีของแถวว่างด้านล่างทั้งหมด
                sws.cell(rr, cc).fill = PatternFill(fill_type=None)

        needed = 4 + len(copied) - 1
        if needed > old_rows:
            sws.insert_rows(old_rows + 1, needed - old_rows)

        for i, (_, row) in enumerate(copied.iterrows(), start=4):
            label = clean_text(row["ปีที่เข้า"])

            if label in ("ป.โท", "ป.เอก"):
                row_fill = red_fill
            elif re.fullmatch(r"\d{4}(?:\.0)?", label):
                row_fill = white_fill
            elif label.startswith("คณะ"):
                row_fill = pink_fill
            else:
                row_fill = green_fill

            # คัดลอกค่าทั้ง 40 คอลัมน์จาก Sheet "สถิติ" โดยตรง
            for c in range(1, 41):
                cell = sws.cell(i, c)
                value = row.iloc[c - 1]
                cell.value = value if pd.notna(value) else None

                # ช่องว่างไม่มีสีพื้น
                if value is None or pd.isna(value):
                    cell.fill = PatternFill(fill_type=None)
                else:
                    cell.fill = row_fill

                cell.font = black_font
                cell.alignment = center
                cell.border = border

            if label in ("ป.โท", "ป.เอก"):
                for c in range(1, 41):
                    sws.cell(i, c).font = Font(
                        name="Tahoma",
                        size=10,
                        color="000000",
                        bold=True,
                    )

            sws.cell(i, 36).number_format = "0.00"
            sws.cell(i, 37).number_format = "0.00"
            sws.cell(i, 40).number_format = "0.00"

        sws.freeze_panes = "A4"
        sws.sheet_view.showGridLines = False
        if len(copied):
            sws.auto_filter.ref = f"A4:AN{needed}"

    # ลบ Sheet เดิมที่แยก ป.โท / ป.เอก
    for old_sheet in ["ป.โท คณะ", "ป.โท ปี", "ป.เอก คณะ", "ป.เอก ปี", "คณะ", "ปี"]:
        if old_sheet in wb.sheetnames:
            del wb[old_sheet]

    # ใช้ข้อมูลจาก Sheet "สถิติ" โดยตรง ไม่คำนวณใหม่
    write_copied_stats_sheet("คณะ", stats, "faculty")
    write_copied_stats_sheet("ปี", stats, "year")

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


st.markdown(
    """
    <style>
    /* พื้นหลังเว็บไซต์สีเทาไล่โทน */
    .stApp {
        background: linear-gradient(135deg, #f5f5f5 0%, #e2e4e7 45%, #c8ccd1 100%);
    }
    .main-title {
        background: linear-gradient(135deg, #0b2f50 0%, #174a73 100%);
        color: white;
        padding: 28px 34px;
        border-radius: 16px;
        margin-bottom: 18px;
        box-shadow: 0 8px 24px rgba(11,47,80,.14);
    }
    .main-title h1 {
        color: white;
        margin: 0;
        font-size: 30px;
        font-weight: 700;
    }
    .main-title p {
        color: #dcecff;
        margin: 8px 0 0 0;
        font-size: 15px;
    }
    .section-title {
        color: #0b2f50;
        font-size: 22px;
        font-weight: 700;
        margin: 22px 0 10px 0;
    }
    .upload-note {
        background: #f3f7fc;
        border-left: 5px solid #2f6690;
        padding: 14px 18px;
        border-radius: 8px;
        color: #29445c;
        margin-bottom: 8px;
    }
    .result-header {
        background: #0b2f50;
        color: white;
        padding: 18px 24px;
        border-radius: 14px 14px 0 0;
        margin-top: 24px;
    }
    .result-header h2 {
        color: white;
        margin: 0;
        font-size: 22px;
    }
    div[data-testid="stFileUploader"] {
        background: white;
        border: 2px dashed #9bb6ce;
        border-radius: 12px;
        padding: 10px;
    }
    div.stButton > button[kind="primary"] {
        background: #f8ad19;
        color: #0b2f50;
        border: none;
        border-radius: 10px;
        min-height: 48px;
        font-size: 17px;
        font-weight: 700;
        box-shadow: 0 4px 12px rgba(248,173,25,.25);
    }
    div.stButton > button[kind="primary"]:hover {
        background: #e99d0d;
        color: #0b2f50;
        border: none;
    }
    div[data-testid="stDownloadButton"] > button {
        background: #f8ad19;
        color: #0b2f50;
        border: none;
        border-radius: 10px;
        min-height: 48px;
        font-size: 17px;
        font-weight: 700;
        box-shadow: 0 4px 12px rgba(248,173,25,.25);
    }
    div[data-testid="stDownloadButton"] > button:hover {
        background: #e99d0d;
        color: #0b2f50;
        border: none;
    }
    </style>

    <div class="main-title">
        <h1>📊 ระบบสถิตินิสิตบัณฑิตศึกษา</h1>
        <p>ระบบประมวลผลข้อมูลนิสิตและจัดทำสถิติอัตโนมัติ</p>
    </div>
    """,
    unsafe_allow_html=True,
)

# ประวัติการประมวลผลสำหรับกราฟแนวโน้มใน Session นี้
if "current_student_history" not in st.session_state:
    st.session_state["current_student_history"] = []

st.markdown('<div class="section-title">📁 นำเข้าข้อมูล</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="upload-note">รองรับไฟล์ Excel ทุกชื่อไฟล์ และระบบจะค้นหา Sheet ที่มีคอลัมน์ข้อมูลนิสิตที่จำเป็นให้อัตโนมัติ</div>',
    unsafe_allow_html=True,
)

uploaded = st.file_uploader(
    "เลือกไฟล์ Excel",
    type=["xlsx", "xls"],
    help="ชื่อไฟล์และชื่อ Sheet ไม่จำเป็นต้องกำหนดตายตัว",
)

if uploaded:
    try:
        df = read_excel(uploaded)
        st.success(f"อ่านข้อมูลสำเร็จ: {len(df):,} รายการ")

        # ============================================================
        # DASHBOARD — ข้อมูลนิสิตปัจจุบัน
        # ============================================================
        # Dashboard ทั้ง 4 ตัวนับสถานะนิสิตปัจจุบัน
        # รวม: นิสิตปัจจุบัน สภาพสมบูรณ์ + รักษาสภาพ + ลาพักการเรียน
        current_statuses = {
            "นิสิตปัจจุบัน สภาพสมบูรณ์",
            "รักษาสภาพนิสิต",
            "ลาพักการเรียน",
        }
        current_df = df[
            df["สถานะนิสิต"].astype(str).str.strip().isin(current_statuses)
        ].copy()

        # ไทย / ต่างชาติ และแยกตามระดับปริญญา
        nationality_norm = current_df["ไทย-ต่างชาติ"].astype(str).str.strip()
        level_norm = current_df["ระดับ"].map(normalize_level)

        thai_mask = nationality_norm == "ไทย"
        foreign_mask = nationality_norm == "ต่างชาติ"
        master_mask = level_norm == "ป.โท"
        doctoral_mask = level_norm == "ป.เอก"

        thai_count = int(thai_mask.sum())
        thai_master_count = int((thai_mask & master_mask).sum())
        thai_doctoral_count = int((thai_mask & doctoral_mask).sum())

        foreign_count = int(foreign_mask.sum())
        foreign_master_count = int((foreign_mask & master_mask).sum())
        foreign_doctoral_count = int((foreign_mask & doctoral_mask).sum())

        total_current_count = int(len(current_df))
        master_total_count = int(master_mask.sum())
        doctoral_total_count = int(doctoral_mask.sum())

        thai_months = [
            "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
            "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"
        ]
        now = datetime.now()
        dashboard_date = f"{now.day} {thai_months[now.month - 1]} {now.year + 543}"

        st.markdown(
            """
            <style>
            .dashboard-card {
                background: #f3f7fc;
                padding: 8px 8px 20px 8px;
                min-height: 120px;
            }
            .dashboard-number {
                color: #0b2a4a;
                font-size: 42px;
                font-weight: 700;
                text-align: center;
                line-height: 1.05;
            }
            .dashboard-label {
                color: #0b2a4a;
                font-size: 16px;
                text-align: center;
                line-height: 1.65;
                margin-top: 16px;
            }
            .dashboard-panel {
                background: #0b2f50;
                color: white;
                padding: 34px 42px 38px 42px;
                margin: 8px 0 0 0;
                min-height: 390px;
            }
            .dashboard-panel h2 {
                color: white;
                font-size: 28px;
                margin: 0 0 10px 0;
            }
            .dashboard-date {
                color: white;
                font-size: 16px;
                font-style: italic;
                margin-bottom: 24px;
            }
            .dashboard-panel ul {
                margin-top: 0;
                padding-left: 22px;
            }
            .dashboard-panel li {
                color: white;
                font-size: 16px;
                line-height: 2.0;
                margin-bottom: 7px;
            }
            </style>
            """,
            unsafe_allow_html=True,
        )

        # Dashboard 4 ตัวเลขด้านบน + กล่องสรุปข้อมูลด้านล่าง
        top_cards = [
            (master_total_count, "จำนวนนิสิต<br>ระดับปริญญาโท"),
            (doctoral_total_count, "จำนวนนิสิต<br>ระดับปริญญาเอก"),
            (thai_count, "จำนวนนิสิต<br>ไทยทั้งหมด"),
            (foreign_count, "จำนวนนิสิต<br>ต่างชาติทั้งหมด"),
        ]

        d1, d2, d3, d4 = st.columns(4)
        for col, (value, label) in zip((d1, d2, d3, d4), top_cards):
            with col:
                st.markdown(
                    f'''
                    <div class="dashboard-card">
                        <div class="dashboard-number">{value:,}</div>
                        <div class="dashboard-label">{label}</div>
                    </div>
                    ''',
                    unsafe_allow_html=True,
                )

        st.markdown(
            f'''
            <div class="dashboard-panel">
                <h2>นิสิตระดับบัณฑิตศึกษา</h2>
                <div class="dashboard-date">ข้อมูล ณ วันที่ {dashboard_date}</div>
                <ul>
                    <li>{total_current_count:,} จำนวนนิสิตทั้งหมด</li>
                    <li>{thai_count:,} จำนวนนิสิตไทยทั้งหมด</li>
                    <li>{thai_master_count:,} จำนวนนิสิตไทย ระดับปริญญาโท</li>
                    <li>{thai_doctoral_count:,} จำนวนนิสิตไทย ระดับปริญญาเอก</li>
                    <li>{foreign_count:,} จำนวนนิสิตต่างชาติทั้งหมด</li>
                    <li>{foreign_master_count:,} จำนวนนิสิตต่างชาติ ระดับปริญญาโท</li>
                    <li>{foreign_doctoral_count:,} จำนวนนิสิตต่างชาติ ระดับปริญญาเอก</li>
                </ul>
            </div>
            ''',
            unsafe_allow_html=True,
        )

        with st.expander("ดูตัวอย่างข้อมูลที่นำเข้า"):
            st.dataframe(df.head(20), use_container_width=True)

        if st.button(
            "🚀 2) ประมวลผลและอัปเดตสถิติ",
            type="primary",
            use_container_width=True,
        ):
            with st.spinner("กำลังประมวลผลข้อมูลและสร้างสถิติ..."):
                processed, stats = build_stats(df)
                excel_bytes = make_excel(
                    uploaded, df, processed, stats
                )

            st.session_state["processed"] = processed
            st.session_state["stats"] = stats
            st.session_state["excel_bytes"] = excel_bytes

            # บันทึกจำนวนนิสิตปัจจุบันของแต่ละครั้งที่กดประมวลผล
            current_statuses_for_history = {
                "นิสิตปัจจุบัน สภาพสมบูรณ์",
                "รักษาสภาพนิสิต",
                "ลาพักการเรียน",
            }
            current_count_for_history = int(
                df["สถานะนิสิต"].astype(str).str.strip().isin(
                    current_statuses_for_history
                ).sum()
            )
            st.session_state["current_student_history"].append(
                {
                    "ครั้งที่": len(st.session_state["current_student_history"]) + 1,
                    "เวลา": datetime.now(),
                    "นิสิตปัจจุบัน": current_count_for_history,
                }
            )

            st.success(
                "ประมวลผลเสร็จแล้ว — สถิติอัปเดตเรียบร้อย "
                "โดยคงสีและรูปแบบของ Sheet สถิติเดิม"
            )

    except Exception as e:
        st.error(f"ไม่สามารถอ่านไฟล์ได้: {e}")


if "stats" in st.session_state:
    stats = st.session_state["stats"]
    processed = st.session_state["processed"]

    st.markdown('<div class="result-header"><h2>📈 ผลสถิติ</h2></div>', unsafe_allow_html=True)

    # ============================================================
    # กราฟแนวโน้มนิสิตปัจจุบันจากแต่ละครั้งที่ประมวลผล
    # ============================================================
    history = pd.DataFrame(st.session_state.get("current_student_history", []))
    if not history.empty:
        st.markdown(
            '<div class="result-header"><h2>📊 แนวโน้มนิสิตปัจจุบัน</h2></div>',
            unsafe_allow_html=True,
        )
        chart_data = history.set_index("ครั้งที่")[["นิสิตปัจจุบัน"]]
        st.line_chart(chart_data, use_container_width=True)

        history_display = history.copy()
        history_display["เวลา"] = history_display["เวลา"].dt.strftime(
            "%d/%m/%Y %H:%M"
        )
        st.dataframe(
            history_display,
            use_container_width=True,
            hide_index=True,
        )

    # ป้องกัน pyarrow/Streamlit ValueError จากชื่อคอลัมน์ซ้ำ
    display_stats = stats.loc[
        :, ~stats.columns.duplicated(keep="last")
    ].copy()

    st.dataframe(
        display_stats,
        use_container_width=True,
        height=600,
    )

    st.download_button(
        "⬇️ 3) ดาวน์โหลด Excel ผลลัพธ์",
        data=st.session_state["excel_bytes"],
        file_name="สถิตินิสิต_ประมวลผลแล้ว.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
    )